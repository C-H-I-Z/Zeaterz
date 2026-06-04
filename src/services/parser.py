"""
parser.py — Document extraction and Gemini parsing service.

Responsibilities:
  1. Extract plain text from uploaded PDF, DOCX, or XLSX files.
  2. Send extracted text to Gemini to identify and structure every regulatory
     requirement found in the document.
  3. Enrich each raw Gemini result with derived fields: year integer, region
     (deterministic prefix lookup), manual-review flag, and metadata.

Public API:
    extract_text(path, filename)      -> str
    parse_with_gemini(text, api_key)  -> (list[dict], token_usage_dict)
    enrich_requirement(req, filename) -> dict
"""

import json
import re
from datetime import datetime, timezone
import pdfplumber
from google import genai
from google.genai import types
from docx import Document
import openpyxl

GEMINI_MODEL = "gemini-2.5-flash"


# ── TEXT EXTRACTION ────────────────────────────────────────────────────────────

def extract_from_pdf(path):
    """Extract all text from a PDF using pdfplumber, concatenated page by page."""
    text = ""

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

    return text


def extract_from_docx(path):
    """Extract text from a DOCX file — body paragraphs first, then table cells.
    
    Tables are flattened to tab-separated rows so Gemini can read structured
    lists that are formatted as tables in Word documents.
    """
    doc = Document(path)
    text = ""

    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text + "\n"

    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                text += row_text + "\n"

    return text


def extract_from_xlsx(path):
    """Flatten all sheets of an XLSX into a tab-separated plain-text string.
    
    data_only=True reads cached cell values rather than formulas.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    text = ""

    for sheet in wb.worksheets:
        text += "\n--- Sheet: " + sheet.title + " ---\n"
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(cell) for cell in row if cell is not None)
            
            if row_text.strip():
                text += row_text + "\n"
    return text


def extract_text(path, filename):
    """Route to the correct extractor based on the uploaded file's extension.
    
    Arguments:
        path:     Absolute path to the temporary file saved on disk.
        filename: Original upload filename — used only to determine the extension.
    
    Returns:
        Extracted plain text as a single string.
    
    Raises:
        ValueError: If the file extension is not pdf, docx, or xlsx.
    """
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return extract_from_pdf(path)
    elif ext == "docx":
        return extract_from_docx(path)
    elif ext == "xlsx":
        return extract_from_xlsx(path)
    
    raise ValueError("Unsupported file type: ." + ext)


# ── REGION INFERENCE ──────────────────────────────────────────────────────────
# Deterministic prefix lookup so region is consistent across runs.
# Gemini's classification is only used as a fallback for unrecognized prefixes.

_US_PREFIXES   = ("21 CFR", "CFR", "FDA", "HIPAA", "HHS", "ANSI", "ASQ", "UL ")
_INTL_PREFIXES = ("ISO", "IEC", "EN ", "EN/", "ICH", "ASTM", "ISTA", "CEN")

def _infer_region(standard_id):
    """Return 'US' or 'International' based on the standard ID prefix, or None if unknown."""
    sid = (standard_id or "").upper().strip()
    for p in _US_PREFIXES:
        if sid.startswith(p.upper()):
            return "US"
        
    for p in _INTL_PREFIXES:
        if sid.startswith(p.upper()):
            return "International"
        
    return None  # unknown — caller falls back to Gemini's value


# ── ENRICHMENT ────────────────────────────────────────────────────────────────

def extract_year(date_str):
    """Pull a 4-digit year (1900–2099) from a date string, or return None.
    
    Returns None for sentinel values (**, Current, blank) so they are
    flagged as needs_manual_review rather than compared against a baseline.
    """
    if not date_str or date_str.strip().lower() in ("current", "**", ""):
        return None
    
    match = re.search(r'\b(19|20)\d{2}\b', str(date_str))

    return int(match.group()) if match else None


def enrich_requirement(req, filename):
    """Normalize and extend one raw Gemini-parsed requirement dict.
    
    Adds derived fields that Gemini doesn't produce directly:
      - date_year:           integer year parsed from the date string
      - needs_manual_review: True when no baseline date exists (**, blank, "Current")
      - region:              deterministic lookup; falls back to Gemini's value
      - source_filename:     original upload name, stored for traceability
      - uploaded_at:         UTC ISO timestamp of this upload session
      - status / current_version / source_url: placeholders filled by checker.py
    
    Args:
        req:      Raw dict from Gemini (standard_id, date, category, region, description).
        filename: Original uploaded filename.
    
    Returns:
        Full requirement dict matching the project data model.
    """
    date_val   = req.get("date", "") or ""
    is_current = date_val.strip().lower() in ("current", "**", "")  # no real date on file
    year       = extract_year(date_val)

    return {
        "standard_id":         req.get("standard_id", ""),
        "date":                date_val,
        "date_year":           int(year) if year is not None else None,
        "category":            req.get("category", ""),
        # Deterministic region lookup overrides Gemini's probabilistic guess.
        "region":              _infer_region(req.get("standard_id", "")) or req.get("region", ""),
        "description":         req.get("description", ""),
        "needs_manual_review": is_current,
        "source_filename":     filename,
        "uploaded_at":         datetime.now(timezone.utc).isoformat(),
        "status":              None,
        "current_version":     None,
        "source_url":          None,
    }


# ── GEMINI PARSING ────────────────────────────────────────────────────────────

def parse_with_gemini(text, api_key):
    """Send extracted document text to Gemini and parse out all regulatory requirements.
    
    Gemini reads the full document text and returns a JSON array where each
    object represents one standard, regulation, or guidance document found.
    Thinking is enabled (default) for better accuracy on complex documents.
    
    Args:
        text:    Plain text extracted from the uploaded file.
        api_key: Gemini API key string.
    
    Returns:
        Tuple of (requirements, tokens) where:
          requirements: list of raw dicts from Gemini (before enrichment)
          tokens:       dict with keys 'prompt', 'output', 'total' (int counts)
    
    Raises:
        json.JSONDecodeError: If Gemini's response cannot be parsed as JSON.
    """
    client = genai.Client(api_key=api_key)

    prompt = (
        "You are a regulatory document parser for medical devices.\n\n"
        "Extract every standard, regulation, or guidance document from the text below.\n"
        "Return ONLY a valid JSON array with no preamble, explanation, or markdown fences.\n\n"
        "Each object must have exactly these fields:\n"
        "- standard_id: the identifier e.g. ISO 13485, 21 CFR Part 820, FDA Guidance 1677\n"
        "- date: the revision or issue date as written. If marked with ** use the word Current\n"
        "- category: the section this belongs to e.g. Quality System Management\n"
        "- region: US or International based on context\n"
        "- description: a brief description of what this standard covers\n\n"
        "Document text:\n" + text
    )

    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=prompt,
    )

    raw = response.text.strip()

    # Strip markdown code fences if Gemini wrapped the JSON in ```json ... ```
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
            
    raw = raw.strip()

    requirements = json.loads(raw)

    # Pull token counts from usage metadata; default to 0 if fields are absent.
    usage  = response.usage_metadata
    tokens = {
        "prompt":  getattr(usage, "prompt_token_count",     0) or 0,
        "output":  getattr(usage, "candidates_token_count", 0) or 0,
        "total":   getattr(usage, "total_token_count",      0) or 0,
    }

    return requirements, tokens