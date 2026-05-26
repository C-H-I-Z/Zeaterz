import json
import re
from datetime import datetime, timezone

import pdfplumber
from google import genai
from google.genai import types
from docx import Document
import openpyxl

GEMINI_MODEL = "gemini-2.5-flash"


def extract_from_pdf(path):
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text


def extract_from_docx(path):
    doc = Document(path)
    text = ""
    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                text += row_text + "\n"
    return text


def extract_from_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    text = ""
    for sheet in wb.worksheets:
        text += "\n--- Sheet: " + sheet.title + " ---\n"
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                text += row_text + "\n"
    return text


def extract_text(path, filename):
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return extract_from_pdf(path)
    elif ext == "docx":
        return extract_from_docx(path)
    elif ext == "xlsx":
        return extract_from_xlsx(path)
    raise ValueError("Unsupported file type: ." + ext)


def extract_year(date_str):
    if not date_str or date_str.strip().lower() in ("current", "**", ""):
        return None
    match = re.search(r'\b(19|20)\d{2}\b', str(date_str))
    return int(match.group()) if match else None


def enrich_requirement(req, filename):
    date_val = req.get("date", "") or ""
    is_current = date_val.strip().lower() in ("current", "**", "")
    year = extract_year(date_val)
    return {
        "standard_id":         req.get("standard_id", ""),
        "date":                date_val,
        "date_year":           int(year) if year is not None else None,
        "category":            req.get("category", ""),
        "region":              req.get("region", ""),
        "description":         req.get("description", ""),
        "needs_manual_review": is_current,
        "source_filename":     filename,
        "uploaded_at":         datetime.now(timezone.utc).isoformat(),
        "status":              None,
        "current_version":     None,
        "source_url":          None,
    }


def parse_with_gemini(text, api_key):
    """Extract requirements from document text via Gemini.

    Returns (list_of_dicts, token_usage_dict).
    """
    client = genai.Client(api_key=api_key)

    prompt = (
        "You are a regulatory document parser for medical devices.\n\n"
        "Extract every standard, regulation, or guidance document from the text below.\n"
        "Return ONLY a valid JSON array with no preamble, explanation, or markdown fences.\n\n"
        "Each object must have exactly these fields:\n"
        "- standard_id: the identifier e.g. ISO 13485, 21 CFR Part 820, FDA Guidance 1677\n"
        "- date: the revision or issue date as written. If marked with ** use the word Current\n"
        "- category: the section this belongs to e.g. Quality System Management\n"
        "- region: US or International based on context\n"
        "- description: a brief description of what this standard covers\n\n"
        "Document text:\n" + text
    )

    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=prompt,
    )

    raw = response.text.strip()
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    raw = raw.strip()

    requirements = json.loads(raw)

    usage = response.usage_metadata
    tokens = {
        "prompt":  getattr(usage, "prompt_token_count",     0) or 0,
        "output":  getattr(usage, "candidates_token_count", 0) or 0,
        "total":   getattr(usage, "total_token_count",      0) or 0,
    }

    return requirements, tokens